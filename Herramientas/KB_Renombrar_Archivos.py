#!/usr/bin/env python3
"""Genera una plantilla Excel desde una carpeta y renombra archivos usando esa plantilla."""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = ["nombre_actual", "nombre_nuevo", "estado", "observacion"]
SHEET_NAME = "Renombrar"


def files_in(folder: Path, recursive: bool, excel_path: Path) -> Iterable[Path]:
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    excluded = {excel_path.resolve(), Path(__file__).resolve()}
    for path in sorted(iterator, key=lambda p: str(p).lower()):
        if path.is_file() and path.resolve() not in excluded and not path.name.startswith("~$"):
            yield path


def relative_name(path: Path, folder: Path) -> str:
    return path.relative_to(folder).as_posix()


def generate_excel(folder: Path, excel_path: Path, recursive: bool) -> None:
    folder = folder.resolve()
    excel_path = excel_path.resolve()
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)

    for path in files_in(folder, recursive, excel_path):
        current = relative_name(path, folder)
        ws.append([current, current, "PENDIENTE", "Editar nombre_nuevo"])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    input_font = Font(color="0000FF")
    imported_font = Font(color="008000")
    caution_fill = PatternFill("solid", fgColor="FCE4D6")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).font = imported_font
        ws.cell(row, 2).font = input_font
        ws.cell(row, 3).fill = caution_fill

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 45

    if ws.max_row >= 2:
        table = Table(displayName="TablaRenombrar", ref=f"A1:D{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)

    instructions = wb.create_sheet("Instrucciones")
    instructions["A1"] = "Cómo usar"
    instructions["A1"].font = Font(size=16, bold=True, color="1F4E78")
    instructions["A3"] = "1. En Renombrar, edita únicamente la columna nombre_nuevo."
    instructions["A4"] = "2. Conserva la extensión correcta, por ejemplo .pdf, .xlsx o .txt."
    instructions["A5"] = "3. Ejecuta el script con el comando renombrar."
    instructions["A6"] = "4. Revisa las columnas estado y observacion después del proceso."
    instructions["A8"] = "Seguridad: el script no sobrescribe archivos existentes y valida nombres duplicados."
    instructions.column_dimensions["A"].width = 105

    wb.save(excel_path)
    print(f"Excel generado: {excel_path}")


def safe_target(folder: Path, value: str) -> Path:
    normalized = value.strip().replace("\\", "/")
    if not normalized:
        raise ValueError("El nombre nuevo está vacío")
    relative = Path(normalized)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("No se permiten rutas absolutas ni '..'")
    target = (folder / relative).resolve()
    try:
        target.relative_to(folder.resolve())
    except ValueError as exc:
        raise ValueError("El destino queda fuera de la carpeta") from exc
    return target


def rename_from_excel(folder: Path, excel_path: Path) -> None:
    folder = folder.resolve()
    excel_path = excel_path.resolve()
    wb = load_workbook(excel_path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"No existe la hoja '{SHEET_NAME}'")
    ws = wb[SHEET_NAME]

    seen_targets: set[str] = set()
    plans: list[tuple[int, Path, Path]] = []

    for row in range(2, ws.max_row + 1):
        current_value = str(ws.cell(row, 1).value or "").strip()
        new_value = str(ws.cell(row, 2).value or "").strip()
        ws.cell(row, 3).value = ""
        ws.cell(row, 4).value = ""

        if not current_value:
            ws.cell(row, 3).value = "OMITIDO"
            ws.cell(row, 4).value = "Sin nombre_actual"
            continue

        source = safe_target(folder, current_value)
        try:
            target = safe_target(folder, new_value)
        except ValueError as exc:
            ws.cell(row, 3).value = "ERROR"
            ws.cell(row, 4).value = str(exc)
            continue

        target_key = os.path.normcase(str(target))
        if target_key in seen_targets:
            ws.cell(row, 3).value = "ERROR"
            ws.cell(row, 4).value = "Destino duplicado en el Excel"
            continue
        seen_targets.add(target_key)

        if source == target:
            ws.cell(row, 3).value = "SIN CAMBIOS"
            ws.cell(row, 4).value = "nombre_actual y nombre_nuevo son iguales"
        elif not source.exists():
            ws.cell(row, 3).value = "ERROR"
            ws.cell(row, 4).value = "El archivo de origen no existe"
        elif target.exists():
            ws.cell(row, 3).value = "ERROR"
            ws.cell(row, 4).value = "Ya existe un archivo con el nombre nuevo"
        else:
            plans.append((row, source, target))

    for row, source, target in plans:
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            source.rename(target)
            ws.cell(row, 1).value = target.relative_to(folder).as_posix()
            ws.cell(row, 3).value = "RENOMBRADO"
            ws.cell(row, 4).value = f"Antes: {source.relative_to(folder).as_posix()}"
        except OSError as exc:
            ws.cell(row, 3).value = "ERROR"
            ws.cell(row, 4).value = str(exc)

    wb.save(excel_path)
    print(f"Proceso terminado. Resultado actualizado en: {excel_path}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Genera Excel y renombra archivos desde Excel")
    sub = parser.add_subparsers(dest="command", required=True)

    gen = sub.add_parser("generar", help="Crear Excel con los archivos de una carpeta")
    gen.add_argument("carpeta", type=Path)
    gen.add_argument("--excel", type=Path, default=Path("plantilla_renombrado.xlsx"))
    gen.add_argument("--recursivo", action="store_true")

    ren = sub.add_parser("renombrar", help="Renombrar según el Excel")
    ren.add_argument("carpeta", type=Path)
    ren.add_argument("--excel", type=Path, default=Path("plantilla_renombrado.xlsx"))
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if args.command == "generar":
            generate_excel(args.carpeta, args.excel, args.recursivo)
        else:
            rename_from_excel(args.carpeta, args.excel)
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
